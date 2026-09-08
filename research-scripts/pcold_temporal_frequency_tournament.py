#!/usr/bin/env python3
from __future__ import annotations

import hashlib
import json
import math
import os
import urllib.request
from collections import Counter
from pathlib import Path

import numpy as np
import openpyxl
import scipy
from scipy.optimize import minimize
from scipy.special import gammaln
from scipy.stats import chi2, kstest, nbinom, poisson

ARTICLE_ID = 23940069
VERSION = 3
FILE_ID = 46360495
WORKBOOK_NAME = "P-COLD-English ver.xlsx"
WORKBOOK_SHA256 = "d3d22f6121b9f405ba3c12a7e0491d907f44a2a3d59d1e8ed105804b5f0e924c"
NOTE = "NOTE：All information that was not disclosed in the original operational risk event text is populated with F"
EXPECTED_EVENTS = 3723
EXPECTED_VALID_DATES = 3047
EXPECTED_EXCLUDED_DATES = 676
EXPECTED_MIN_MONTH = (1988, 6)
EXPECTED_MAX_MONTH = (2022, 11)
TRAIN_END = (2014, 12)
VALID_START = (2015, 1)
VALID_END = (2018, 12)
FINAL_START = (2019, 1)
BLOCK_LENGTH = 6
BOOT_REPS = 10_000
SEED = 20260907
TAIL_P = 0.995
SEVERITY = 100.0


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def integer(v):
    try:
        f = float(v)
        if f.is_integer():
            return int(f)
    except Exception:
        pass
    return None


def month_key(y: int, m: int) -> int:
    return y * 12 + (m - 1)


def month_pair(k: int) -> tuple[int, int]:
    return k // 12, k % 12 + 1


def fetch_workbook() -> bytes:
    api = f"https://api.figshare.com/v2/articles/{ARTICLE_ID}/versions/{VERSION}"
    req = urllib.request.Request(api, headers={"User-Agent": "OpenCatastrophe-data P-COLD temporal tournament"})
    with urllib.request.urlopen(req, timeout=60) as r:
        meta = json.loads(r.read())
    if meta.get("version") != VERSION:
        raise RuntimeError(f"FAIL_CLOSED Figshare version {meta.get('version')!r}")
    if str(meta.get("doi", "")).lower() != "10.6084/m9.figshare.23940069.v3":
        raise RuntimeError(f"FAIL_CLOSED DOI {meta.get('doi')!r}")
    files = [f for f in meta.get("files", []) if f.get("id") == FILE_ID and f.get("name") == WORKBOOK_NAME]
    if len(files) != 1:
        raise RuntimeError("FAIL_CLOSED exact workbook identity")
    req = urllib.request.Request(files[0]["download_url"], headers={"User-Agent": "OpenCatastrophe-data P-COLD temporal tournament"})
    with urllib.request.urlopen(req, timeout=120) as r:
        data = r.read()
    if sha256_bytes(data) != WORKBOOK_SHA256:
        raise RuntimeError("FAIL_CLOSED workbook SHA256")
    if files[0].get("size") is not None and int(files[0]["size"]) != len(data):
        raise RuntimeError("FAIL_CLOSED workbook byte length")
    return data


def parse_exact_dates(data: bytes):
    p = Path("pcold-v3.xlsx")
    p.write_bytes(data)
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    if wb.sheetnames != ["Sheet1"]:
        raise RuntimeError(f"FAIL_CLOSED sheet identity {wb.sheetnames!r}")
    ws = wb["Sheet1"]
    rows = ws.iter_rows(values_only=True)
    hdr = [str(x).strip() if x is not None else "" for x in next(rows)]
    ix = {h: hdr.index(h) for h in ["Event ID", "Occurrence year", "Occurrence month"]}
    ids = []
    valid = []
    excluded = []
    non_event = []
    for rn, row in enumerate(rows, start=2):
        eid_raw = row[ix["Event ID"]] if ix["Event ID"] < len(row) else None
        if eid_raw in (None, ""):
            continue
        eid = integer(eid_raw)
        if eid is None:
            non_event.append((rn, str(eid_raw).strip()))
            continue
        ids.append(eid)
        yraw = row[ix["Occurrence year"]] if ix["Occurrence year"] < len(row) else None
        mraw = row[ix["Occurrence month"]] if ix["Occurrence month"] < len(row) else None
        y = integer(yraw)
        m = integer(mraw)
        if y is None:
            reason = "occurrence_year_undisclosed_F" if str(yraw).strip().upper() == "F" else "occurrence_year_unparseable"
        elif m is None:
            reason = "occurrence_month_undisclosed_F" if str(mraw).strip().upper() == "F" else "occurrence_month_unparseable"
        elif not 1 <= m <= 12:
            reason = "occurrence_month_outside_1_12"
        else:
            reason = None
        if reason is None:
            valid.append((eid, y, m))
        else:
            if y is None:
                part = "unknown_year"
            elif (y, 12) <= TRAIN_END:
                part = "training"
            elif y <= VALID_END[0]:
                part = "validation"
            else:
                part = "final"
            excluded.append((eid, reason, part))
    if sorted(ids) != list(range(1, EXPECTED_EVENTS + 1)):
        raise RuntimeError("FAIL_CLOSED Event IDs not exactly 1..3723")
    if non_event != [(3726, NOTE)]:
        raise RuntimeError(f"FAIL_CLOSED non-event note changed: {non_event!r}")
    if len(valid) != EXPECTED_VALID_DATES or len(excluded) != EXPECTED_EXCLUDED_DATES:
        raise RuntimeError(f"FAIL_CLOSED exact-date counts valid={len(valid)} excluded={len(excluded)}")
    minm = min((y, m) for _, y, m in valid)
    maxm = max((y, m) for _, y, m in valid)
    if minm != EXPECTED_MIN_MONTH or maxm != EXPECTED_MAX_MONTH:
        raise RuntimeError(f"FAIL_CLOSED exact-date span {minm}..{maxm}")
    return valid, excluded


def monthly_series(valid):
    lo = month_key(*EXPECTED_MIN_MONTH)
    hi = month_key(*EXPECTED_MAX_MONTH)
    keys = np.arange(lo, hi + 1, dtype=int)
    counts = Counter(month_key(y, m) for _, y, m in valid)
    y = np.array([counts[int(k)] for k in keys], dtype=float)
    pairs = [month_pair(int(k)) for k in keys]
    return keys, pairs, y


def poisson_logpmf(y: np.ndarray, mu: np.ndarray) -> np.ndarray:
    mu = np.clip(np.asarray(mu, dtype=float), 1e-12, 1e12)
    return y * np.log(mu) - mu - gammaln(y + 1.0)


def nb2_logpmf(y: np.ndarray, mu: np.ndarray, alpha: float) -> np.ndarray:
    if alpha <= 1e-10:
        return poisson_logpmf(y, mu)
    mu = np.clip(np.asarray(mu, dtype=float), 1e-12, 1e12)
    r = 1.0 / alpha
    p = r / (r + mu)
    return gammaln(y + r) - gammaln(r) - gammaln(y + 1.0) + r * np.log(p) + y * np.log1p(-p)


def design(keys: np.ndarray, pairs, train_mask: np.ndarray) -> np.ndarray:
    idx = np.arange(len(keys), dtype=float)
    center = float(idx[train_mask].mean())
    trend_years = (idx - center) / 12.0
    months = np.array([m for _, m in pairs], dtype=float)
    return np.column_stack(
        [
            np.ones(len(keys)),
            trend_years,
            np.sin(2.0 * np.pi * months / 12.0),
            np.cos(2.0 * np.pi * months / 12.0),
        ]
    )


def fit_models(y: np.ndarray, X: np.ndarray, train: np.ndarray):
    yt = y[train]
    Xt = X[train]
    mu0 = max(float(yt.mean()), 1e-12)
    nll0 = float(-poisson_logpmf(yt, np.full(len(yt), mu0)).sum())

    def f1(z):
        mu = math.exp(float(z[0])); alpha = math.exp(float(z[1]))
        return float(-nb2_logpmf(yt, np.full(len(yt), mu), alpha).sum())

    r1 = minimize(f1, np.array([math.log(mu0), math.log(0.1)]), method="L-BFGS-B", bounds=[(-20, 20), (-20, 8)])
    if not r1.success or not np.isfinite(r1.fun):
        raise RuntimeError(f"M1 optimizer failed: {r1.message}")
    if float(r1.fun) + 1e-9 < nll0:
        m1_mu = math.exp(float(r1.x[0])); m1_alpha = math.exp(float(r1.x[1])); m1_boundary = False
    else:
        m1_mu = mu0; m1_alpha = 0.0; m1_boundary = True

    def f2(beta):
        mu = np.exp(np.clip(Xt @ beta, -20.0, 20.0))
        return float(-poisson_logpmf(yt, mu).sum())

    b0 = np.array([math.log(mu0), 0.0, 0.0, 0.0])
    r2 = minimize(f2, b0, method="L-BFGS-B", bounds=[(-20, 20)] * 4)
    if not r2.success or not np.isfinite(r2.fun):
        raise RuntimeError(f"M2 optimizer failed: {r2.message}")
    beta2 = np.array(r2.x, dtype=float)
    nll2 = float(r2.fun)

    init_alpha = m1_alpha if m1_alpha > 1e-8 else 0.1
    def f3(z):
        beta = z[:4]; alpha = math.exp(float(z[4]))
        mu = np.exp(np.clip(Xt @ beta, -20.0, 20.0))
        return float(-nb2_logpmf(yt, mu, alpha).sum())

    r3 = minimize(
        f3,
        np.r_[beta2, math.log(init_alpha)],
        method="L-BFGS-B",
        bounds=[(-20, 20)] * 4 + [(-20, 8)],
    )
    if not r3.success or not np.isfinite(r3.fun):
        raise RuntimeError(f"M3 optimizer failed: {r3.message}")
    if float(r3.fun) + 1e-9 < nll2:
        beta3 = np.array(r3.x[:4], dtype=float); alpha3 = math.exp(float(r3.x[4])); m3_boundary = False
    else:
        beta3 = beta2.copy(); alpha3 = 0.0; m3_boundary = True

    return {
        "M0": {"family": "poisson_stationary", "mu": mu0, "alpha": 0.0, "converged": True},
        "M1": {"family": "nb2_stationary", "mu": m1_mu, "alpha": m1_alpha, "poisson_boundary": m1_boundary, "optimizer": str(r1.message), "converged": bool(r1.success)},
        "M2": {"family": "poisson_calendar", "beta": beta2.tolist(), "alpha": 0.0, "optimizer": str(r2.message), "converged": bool(r2.success)},
        "M3": {"family": "nb2_calendar", "beta": beta3.tolist(), "alpha": alpha3, "poisson_boundary": m3_boundary, "optimizer": str(r3.message), "converged": bool(r3.success)},
    }


def predictive_mu(model, X):
    fam = model["family"]
    if fam.endswith("stationary"):
        return np.full(len(X), float(model["mu"]))
    return np.exp(np.clip(X @ np.array(model["beta"], dtype=float), -20.0, 20.0))


def logpmf_model(y, model, X):
    mu = predictive_mu(model, X)
    alpha = float(model.get("alpha", 0.0))
    return nb2_logpmf(y, mu, alpha) if alpha > 1e-10 else poisson_logpmf(y, mu)


def scipy_dist(model, mu: float):
    alpha = float(model.get("alpha", 0.0))
    if alpha <= 1e-10:
        return poisson(mu)
    r = 1.0 / alpha
    p = r / (r + mu)
    return nbinom(r, p)


def diagnostics(y, model, X, seed):
    mu = predictive_mu(model, X)
    alpha = float(model.get("alpha", 0.0))
    rng = np.random.default_rng(seed)
    u = rng.random(len(y))
    pits = np.empty(len(y))
    cover = {"50": 0, "90": 0, "95": 0}
    quantiles = {"50": (0.25, 0.75), "90": (0.05, 0.95), "95": (0.025, 0.975)}
    for i, (obs, mm) in enumerate(zip(y.astype(int), mu)):
        d = scipy_dist(model, float(mm))
        prev = 0.0 if obs <= 0 else float(d.cdf(obs - 1))
        mass = float(d.pmf(obs))
        pits[i] = prev + u[i] * mass
        for k, (a, b) in quantiles.items():
            lo = int(d.ppf(a)); hi = int(d.ppf(b))
            if lo <= obs <= hi:
                cover[k] += 1
    var = mu + alpha * mu * mu
    resid = (y - mu) / np.sqrt(np.clip(var, 1e-12, None))
    def acf(lag):
        if len(resid) <= lag:
            return None
        a = resid[:-lag] - resid[:-lag].mean(); b = resid[lag:] - resid[lag:].mean()
        den = math.sqrt(float(np.dot(a, a) * np.dot(b, b)))
        return None if den == 0 else float(np.dot(a, b) / den)
    acfs = {}
    for lag in range(1, 13):
        acfs[lag] = acf(lag)
    valid_rhos = [acfs[k] for k in range(1, 13)]
    if all(v is not None and math.isfinite(v) for v in valid_rhos) and len(resid) > 12:
        q = len(resid) * (len(resid) + 2.0) * sum(float(acfs[k]) ** 2 / (len(resid) - k) for k in range(1, 13))
        lb = {"Q12": float(q), "pvalue_chi2_df12": float(chi2.sf(q, 12))}
    else:
        lb = None
    ks = kstest(pits, "uniform")
    cov = {k: cover[k] / len(y) for k in cover}
    catastrophic = not (
        0.20 <= cov["50"] <= 0.80
        and cov["90"] >= 0.70
        and cov["95"] >= 0.80
        and float(ks.statistic) <= 0.30
    )
    return {
        "mean_nll": float(-logpmf_model(y, model, X).mean()),
        "randomized_pit": {"seed": seed, "mean": float(pits.mean()), "ks_statistic": float(ks.statistic), "ks_pvalue": float(ks.pvalue)},
        "central_interval_coverage": cov,
        "pearson_residual": {"mean": float(resid.mean()), "variance_population": float(resid.var(ddof=0)), "acf_lag1": acfs[1], "acf_lag12": acfs[12], "ljung_box": lb},
        "catastrophic_calibration_failure": catastrophic,
        "catastrophic_rule": "50% coverage must be in [0.20,0.80], 90% >=0.70, 95% >=0.80, randomized-PIT KS statistic <=0.30",
    }


def moving_block_ci(improvement: np.ndarray, seed: int):
    x = np.asarray(improvement, dtype=float)
    n = len(x)
    L = min(BLOCK_LENGTH, n)
    starts_n = n - L + 1
    blocks = math.ceil(n / L)
    rng = np.random.default_rng(seed)
    vals = np.empty(BOOT_REPS)
    for b in range(BOOT_REPS):
        starts = rng.integers(0, starts_n, size=blocks)
        idx = np.concatenate([np.arange(s, s + L) for s in starts])[:n]
        vals[b] = float(x[idx].mean())
    return {
        "block_length_months": L,
        "reps": BOOT_REPS,
        "seed": seed,
        "mean_improvement": float(x.mean()),
        "q025": float(np.quantile(vals, 0.025)),
        "median": float(np.quantile(vals, 0.5)),
        "q975": float(np.quantile(vals, 0.975)),
        "resampling": "ordinary moving blocks: sample contiguous length-6 start positions with replacement, concatenate and truncate to partition length",
    }


def partition_receipt(y, X, models, mask, label):
    yy = y[mask]; XX = X[mask]
    model_diag = {name: diagnostics(yy, model, XX, SEED) for name, model in models.items()}
    loss = {name: -logpmf_model(yy, model, XX) for name, model in models.items()}
    comparisons = {}
    for name in ["M1", "M2", "M3"]:
        comparisons[f"{name}-M0"] = moving_block_ci(loss["M0"] - loss[name], SEED)
    comparisons["M3-M2"] = moving_block_ci(loss["M2"] - loss["M3"], SEED)
    return {
        "label": label,
        "months": int(len(yy)),
        "raw_count_mean": float(yy.mean()),
        "raw_count_variance_population": float(yy.var(ddof=0)),
        "models": model_diag,
        "paired_nll_improvement": comparisons,
    }


def validation_selection(validation):
    passed = []
    for name in ["M1", "M2", "M3"]:
        comp = validation["paired_nll_improvement"][f"{name}-M0"]
        diag = validation["models"][name]
        ok = (
            comp["mean_improvement"] > 0
            and comp["q025"] > 0
            and not diag["catastrophic_calibration_failure"]
        )
        if ok:
            passed.append(name)
    if not passed:
        return {"passed_challengers": [], "winner": "M0", "tie_within_1e-6": False}
    nlls = {name: validation["models"][name]["mean_nll"] for name in passed}
    best = min(nlls.values())
    tied = sorted([name for name, v in nlls.items() if abs(v - best) <= 1e-6])
    if len(tied) > 1:
        winner = "MIXED_TIE_NO_COMPLEXITY_PROMOTION"
        tie = True
    else:
        winner = tied[0]
        tie = False
    return {"passed_challengers": passed, "winner": winner, "tie_within_1e-6": tie, "challenger_mean_nll": nlls}


def monthly_pmf(model, mu: float, tol=1e-12):
    d = scipy_dist(model, mu)
    hi = int(d.ppf(1.0 - tol))
    if not np.isfinite(hi) or hi < 0:
        raise RuntimeError("FAIL_CLOSED non-finite tail support")
    xs = np.arange(hi + 1)
    p = np.asarray(d.pmf(xs), dtype=float)
    if p.sum() <= 0:
        raise RuntimeError("FAIL_CLOSED zero pmf mass")
    return p


def convolve_year(model, Xyear):
    mus = predictive_mu(model, Xyear)
    total = np.array([1.0])
    for mu in mus:
        total = np.convolve(total, monthly_pmf(model, float(mu)))
    total = np.clip(total, 0.0, None)
    total /= total.sum()
    return total


def discrete_tail_metrics(pmf: np.ndarray, attachment_loss: float):
    cdf = np.cumsum(pmf)
    k = int(np.searchsorted(cdf, TAIL_P, side="left"))
    f_at = float(cdf[k])
    tail_above = float(np.dot(np.arange(k + 1, len(pmf)), pmf[k + 1:]))
    es_count = (tail_above + k * (f_at - TAIL_P)) / (1.0 - TAIL_P)
    xs_loss = SEVERITY * np.arange(len(pmf), dtype=float)
    stop = float(np.dot(np.maximum(xs_loss - attachment_loss, 0.0), pmf))
    return {
        "VaR99.5": float(SEVERITY * k),
        "ES99.5_quantile_integral": float(SEVERITY * es_count),
        "ES_definition": "(integral_p^1 VaR_u du)/(1-p), implemented exactly for discrete mass including only F(VaR)-p of the VaR atom",
        "stop_loss_expectation": stop,
        "pmf_mass_after_normalization": float(pmf.sum()),
    }


def tail_sensitivity(pairs, X, models, train_mask):
    train_mean = None
    # The fixed attachment is based on observed training monthly mean count, supplied by caller later.
    years = sorted(set(y for y, m in pairs if (y, m) >= FINAL_START))
    complete = []
    for year in years:
        idx = [i for i, (yy, mm) in enumerate(pairs) if yy == year]
        months = sorted(pairs[i][1] for i in idx)
        if months == list(range(1, 13)):
            complete.append((year, np.array(idx, dtype=int)))
    return complete


def main():
    data = fetch_workbook()
    valid, excluded = parse_exact_dates(data)
    keys, pairs, y = monthly_series(valid)
    X = design(keys, pairs, np.array([(ym <= TRAIN_END) for ym in pairs], dtype=bool))
    train = np.array([ym <= TRAIN_END for ym in pairs], dtype=bool)
    validation = np.array([VALID_START <= ym <= VALID_END for ym in pairs], dtype=bool)
    final = np.array([ym >= FINAL_START for ym in pairs], dtype=bool)
    if not train.any() or validation.sum() != 48 or not final.any():
        raise RuntimeError(f"FAIL_CLOSED partition sizes train={train.sum()} validation={validation.sum()} final={final.sum()}")
    if pairs[final][0] if False else False:
        pass
    models = fit_models(y, X, train)
    val_receipt = partition_receipt(y, X, models, validation, "validation_2015_2018")
    selection = validation_selection(val_receipt)
    final_receipt = partition_receipt(y, X, models, final, "final_2019_to_2022_11_exact_date_admissible")

    train_mean = float(y[train].mean())
    attachment = 2.0 * SEVERITY * (12.0 * train_mean)
    complete_years = tail_sensitivity(pairs, X, models, train)
    annual = {}
    for year, idx in complete_years:
        annual[str(year)] = {}
        for name, model in models.items():
            pmf = convolve_year(model, X[idx])
            annual[str(year)][name] = discrete_tail_metrics(pmf, attachment)
    final_support = {}
    for name in ["M1", "M2", "M3"]:
        c = final_receipt["paired_nll_improvement"][f"{name}-M0"]
        final_support[name] = bool(c["mean_improvement"] > 0 and c["q025"] > 0)
    m3m2 = final_receipt["paired_nll_improvement"]["M3-M2"]
    if m3m2["mean_improvement"] > 0 and m3m2["q025"] > 0 and float(models["M3"]["alpha"]) > 1e-10:
        pattern = "RESIDUAL_OVERDISPERSION_SUPPORTED_IN_FINAL"
    elif final_support["M2"] and not (m3m2["mean_improvement"] > 0 and m3m2["q025"] > 0):
        pattern = "CALENDAR_STRUCTURE_SUPPORTED_RESIDUAL_NB_NOT_RESOLVED"
    elif not any(final_support.values()):
        pattern = "NO_CHALLENGER_RESOLVED_FINAL"
    else:
        pattern = "MIXED"

    ex_reason = Counter(reason for _, reason, _ in excluded)
    ex_part = Counter(part for _, _, part in excluded)
    receipt = {
        "schema": "pcold-temporal-frequency-tournament-v0",
        "status": "EXECUTED_ONCE_AFTER_SOURCE_LOCK",
        "ffbk_preregistration_pr": 1462,
        "source": {
            "doi": "10.6084/m9.figshare.23940069.v3",
            "article_id": ARTICLE_ID,
            "version": VERSION,
            "file_id": FILE_ID,
            "filename": WORKBOOK_NAME,
            "sha256": WORKBOOK_SHA256,
            "event_ids_exact_1_to_3723": True,
            "valid_exact_occurrence_month_rows": len(valid),
            "excluded_occurrence_date_rows": len(excluded),
            "excluded_by_reason": dict(ex_reason),
            "excluded_partition_membership_from_year_if_known": dict(ex_part),
            "monthly_series_start": f"{pairs[0][0]:04d}-{pairs[0][1]:02d}",
            "monthly_series_end": f"{pairs[-1][0]:04d}-{pairs[-1][1]:02d}",
            "endpoint_boundary": "Exact-date admissible series ends 2022-11. No synthetic zero months are appended into 2023 because undisclosed-month events exist; 2023 is not treated as observed zero-count time.",
        },
        "environment": {
            "python": os.sys.version,
            "numpy": np.__version__,
            "scipy": scipy.__version__,
            "openpyxl": openpyxl.__version__,
            "platform": os.uname().sysname + " " + os.uname().release,
        },
        "execution": {
            "repository": os.environ.get("GITHUB_REPOSITORY"),
            "workflow": os.environ.get("GITHUB_WORKFLOW"),
            "run_id": os.environ.get("GITHUB_RUN_ID"),
            "run_attempt": os.environ.get("GITHUB_RUN_ATTEMPT"),
            "head_sha": os.environ.get("GITHUB_SHA"),
            "bootstrap_seed": SEED,
            "bootstrap_reps": BOOT_REPS,
            "block_length_months": BLOCK_LENGTH,
            "pit_seed": SEED,
        },
        "partition": {
            "training_months": int(train.sum()),
            "validation_months": int(validation.sum()),
            "final_months": int(final.sum()),
            "training_end": "2014-12",
            "validation": "2015-01..2018-12",
            "final": "2019-01..2022-11 exact-date admissible endpoint",
        },
        "training": {
            "raw_count_mean": train_mean,
            "raw_count_variance_population": float(y[train].var(ddof=0)),
            "models": models,
        },
        "validation": val_receipt,
        "validation_selection": selection,
        "final": final_receipt,
        "final_supported_vs_M0": final_support,
        "interpretive_pattern": pattern,
        "secondary_frequency_only_tail": {
            "severity_per_event": SEVERITY,
            "parameter_uncertainty": "omitted by preregistration",
            "independence": "monthly predictive counts are convolved as conditionally independent under each frozen plug-in model",
            "stop_loss_attachment": attachment,
            "attachment_formula": "2 * 100 * (12 * observed training monthly mean count)",
            "complete_final_calendar_years": [year for year, _ in complete_years],
            "annual": annual,
        },
        "scientific_boundary": [
            "external public-media Chinese banking benchmark only",
            "exact-date admissible subset; 676 events with undisclosed/unparseable occurrence date are excluded under the preregistered source-quality rule",
            "no causal interpretation of calendar trend",
            "no insurer/P&C default transfer",
            "no parameter uncertainty in secondary tail sensitivity",
            "no production, regulatory, capital, schema, or NextGen API promotion",
        ],
    }
    out = Path("research-results/pcold-temporal-frequency-tournament-v0.json")
    out.write_text(json.dumps(receipt, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print("PCOLD_TEMPORAL_TOURNAMENT_OK")
    print(json.dumps({
        "validation_selection": selection,
        "final_supported_vs_M0": final_support,
        "interpretive_pattern": pattern,
        "M1_alpha": models["M1"]["alpha"],
        "M3_alpha": models["M3"]["alpha"],
        "validation_mean_nll": {k:v["mean_nll"] for k,v in val_receipt["models"].items()},
        "final_mean_nll": {k:v["mean_nll"] for k,v in final_receipt["models"].items()},
        "complete_final_years": [year for year, _ in complete_years],
    }, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
