#!/usr/bin/env python3
from __future__ import annotations

import hashlib
import json
import math
import platform
import urllib.request
from collections import Counter
from pathlib import Path

import numpy as np
import openpyxl
import scipy
from scipy.optimize import minimize
from scipy.special import expit, gammaln
from scipy.stats import chi2, kstest, nbinom, poisson

ARTICLE_ID = 23940069
VERSION = 3
FILE_ID = 46360495
WORKBOOK_NAME = "P-COLD-English ver.xlsx"
WORKBOOK_SHA256 = "d3d22f6121b9f405ba3c12a7e0491d907f44a2a3d59d1e8ed105804b5f0e924c"
NOTE = "NOTE：All information that was not disclosed in the original operational risk event text is populated with F"
EXPECTED_EVENTS = 3723
EXPECTED_VALID_DATES = 3047
EXPECTED_EXCLUDED_DATES = 676
EXPECTED_MIN_MONTH = (1988, 6)
EXPECTED_MAX_MONTH = (2022, 11)
SCORE_START = (2005, 1)
PRIMARY_BLOCK = 12
SECONDARY_BLOCK = 6
BOOT_REPS = 10_000
SEED = 20260908
PERSISTENCE_CAP = 0.995

MODEL_NAMES = ("S0", "S1", "D0", "D1")


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def integer(v):
    try:
        f = float(v)
        if f.is_integer():
            return int(f)
    except Exception:
        pass
    return None


def month_key(y: int, m: int) -> int:
    return y * 12 + (m - 1)


def month_pair(k: int) -> tuple[int, int]:
    return k // 12, k % 12 + 1


def fetch_workbook() -> bytes:
    api = f"https://api.figshare.com/v2/articles/{ARTICLE_ID}/versions/{VERSION}"
    req = urllib.request.Request(api, headers={"User-Agent": "OpenCatastrophe-data P-COLD INGARCH prequential"})
    with urllib.request.urlopen(req, timeout=60) as r:
        meta = json.loads(r.read())
    if meta.get("version") != VERSION:
        raise RuntimeError(f"FAIL_CLOSED Figshare version {meta.get('version')!r}")
    if str(meta.get("doi", "")).lower() != "10.6084/m9.figshare.23940069.v3":
        raise RuntimeError(f"FAIL_CLOSED DOI {meta.get('doi')!r}")
    files = [f for f in meta.get("files", []) if f.get("id") == FILE_ID and f.get("name") == WORKBOOK_NAME]
    if len(files) != 1:
        raise RuntimeError("FAIL_CLOSED exact workbook identity")
    req = urllib.request.Request(files[0]["download_url"], headers={"User-Agent": "OpenCatastrophe-data P-COLD INGARCH prequential"})
    with urllib.request.urlopen(req, timeout=120) as r:
        data = r.read()
    if sha256_bytes(data) != WORKBOOK_SHA256:
        raise RuntimeError("FAIL_CLOSED workbook SHA256")
    if files[0].get("size") is not None and int(files[0]["size"]) != len(data):
        raise RuntimeError("FAIL_CLOSED workbook byte length")
    return data


def parse_exact_dates(data: bytes):
    p = Path("pcold-v3.xlsx")
    p.write_bytes(data)
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    if wb.sheetnames != ["Sheet1"]:
        raise RuntimeError(f"FAIL_CLOSED sheet identity {wb.sheetnames!r}")
    ws = wb["Sheet1"]
    rows = ws.iter_rows(values_only=True)
    hdr = [str(x).strip() if x is not None else "" for x in next(rows)]
    ix = {h: hdr.index(h) for h in ["Event ID", "Occurrence year", "Occurrence month"]}
    ids = []
    valid = []
    excluded = []
    non_event = []
    for rn, row in enumerate(rows, start=2):
        eid_raw = row[ix["Event ID"]] if ix["Event ID"] < len(row) else None
        if eid_raw in (None, ""):
            continue
        eid = integer(eid_raw)
        if eid is None:
            non_event.append((rn, str(eid_raw).strip()))
            continue
        ids.append(eid)
        yraw = row[ix["Occurrence year"]] if ix["Occurrence year"] < len(row) else None
        mraw = row[ix["Occurrence month"]] if ix["Occurrence month"] < len(row) else None
        y = integer(yraw)
        m = integer(mraw)
        if y is None:
            reason = "occurrence_year_undisclosed_F" if str(yraw).strip().upper() == "F" else "occurrence_year_unparseable"
        elif m is None:
            reason = "occurrence_month_undisclosed_F" if str(mraw).strip().upper() == "F" else "occurrence_month_unparseable"
        elif not 1 <= m <= 12:
            reason = "occurrence_month_outside_1_12"
        else:
            reason = None
        if reason is None:
            valid.append((eid, y, m))
        else:
            excluded.append((eid, reason))
    if sorted(ids) != list(range(1, EXPECTED_EVENTS + 1)):
        raise RuntimeError("FAIL_CLOSED Event IDs not exactly 1..3723")
    if non_event != [(3726, NOTE)]:
        raise RuntimeError(f"FAIL_CLOSED non-event note changed: {non_event!r}")
    if len(valid) != EXPECTED_VALID_DATES or len(excluded) != EXPECTED_EXCLUDED_DATES:
        raise RuntimeError(f"FAIL_CLOSED exact-date counts valid={len(valid)} excluded={len(excluded)}")
    minm = min((y, m) for _, y, m in valid)
    maxm = max((y, m) for _, y, m in valid)
    if minm != EXPECTED_MIN_MONTH or maxm != EXPECTED_MAX_MONTH:
        raise RuntimeError(f"FAIL_CLOSED exact-date span {minm}..{maxm}")
    return valid, excluded


def monthly_series(valid):
    lo = month_key(*EXPECTED_MIN_MONTH)
    hi = month_key(*EXPECTED_MAX_MONTH)
    keys = np.arange(lo, hi + 1, dtype=int)
    counts = Counter(month_key(y, m) for _, y, m in valid)
    y = np.array([counts[int(k)] for k in keys], dtype=float)
    pairs = [month_pair(int(k)) for k in keys]
    return keys, pairs, y


def poisson_logpmf(y: np.ndarray, mu: np.ndarray) -> np.ndarray:
    y = np.asarray(y, dtype=float)
    mu = np.clip(np.asarray(mu, dtype=float), 1e-12, 1e12)
    return y * np.log(mu) - mu - gammaln(y + 1.0)


def nb2_logpmf(y: np.ndarray, mu: np.ndarray, kappa: float) -> np.ndarray:
    y = np.asarray(y, dtype=float)
    mu = np.clip(np.asarray(mu, dtype=float), 1e-12, 1e12)
    if kappa <= 1e-10:
        return poisson_logpmf(y, mu)
    r = 1.0 / float(kappa)
    return (
        gammaln(y + r)
        - gammaln(r)
        - gammaln(y + 1.0)
        - r * np.log1p(mu / r)
        + y * (np.log(mu) - np.log(r + mu))
    )


def logit(p: float) -> float:
    p = min(max(float(p), 1e-12), 1.0 - 1e-12)
    return math.log(p / (1.0 - p))


def ingarch_transform(theta: np.ndarray) -> tuple[float, float, float]:
    omega = math.exp(float(theta[0]))
    s = PERSISTENCE_CAP * float(expit(theta[1]))
    q = float(expit(theta[2]))
    a = s * q
    b = s * (1.0 - q)
    return omega, a, b


def ingarch_start(history_mean: float, a: float, b: float) -> np.ndarray:
    s = a + b
    if not 0.0 < s < PERSISTENCE_CAP:
        raise ValueError("start persistence outside frozen domain")
    q = a / s
    omega = max(float(history_mean) * (1.0 - s), 1e-8)
    return np.array([math.log(omega), logit(s / PERSISTENCE_CAP), logit(q)], dtype=float)


def ingarch_lambda(history: np.ndarray, omega: float, a: float, b: float) -> np.ndarray:
    history = np.asarray(history, dtype=float)
    if len(history) == 0:
        return np.array([], dtype=float)
    if not (omega > 0 and a >= 0 and b >= 0 and a + b <= PERSISTENCE_CAP + 1e-12):
        raise ValueError("invalid INGARCH parameters")
    denom = max(1.0 - a - b, 1e-12)
    lam = np.empty(len(history), dtype=float)
    lam[0] = omega / denom
    for i in range(1, len(history)):
        lam[i] = omega + a * history[i - 1] + b * lam[i - 1]
    return np.clip(lam, 1e-12, 1e12)


def ingarch_forecast(history: np.ndarray, omega: float, a: float, b: float) -> float:
    history = np.asarray(history, dtype=float)
    if len(history) == 0:
        return omega / max(1.0 - a - b, 1e-12)
    lam = ingarch_lambda(history, omega, a, b)
    return float(np.clip(omega + a * history[-1] + b * lam[-1], 1e-12, 1e12))


def fit_s0(history: np.ndarray) -> dict:
    mu = max(float(np.mean(history)), 1e-12)
    return {"family": "poisson_iid", "mu": mu, "kappa": 0.0, "converged": True}


def fit_s1(history: np.ndarray) -> dict:
    history = np.asarray(history, dtype=float)
    mu0 = max(float(history.mean()), 1e-12)
    poiss_nll = float(-poisson_logpmf(history, np.full(len(history), mu0)).sum())

    starts = [0.1, 1.0, max(float(history.var() - history.mean()) / max(history.mean() ** 2, 1e-12), 1e-4)]
    best = None
    for kap0 in starts:
        x0 = np.array([math.log(mu0), math.log(min(max(kap0, math.exp(-20)), math.exp(6)))])

        def obj(z):
            mu = math.exp(float(z[0]))
            kap = math.exp(float(z[1]))
            return float(-nb2_logpmf(history, np.full(len(history), mu), kap).sum())

        r = minimize(obj, x0, method="L-BFGS-B", bounds=[(-20, 20), (-20, 6)])
        if r.success and np.isfinite(r.fun) and (best is None or r.fun < best.fun):
            best = r
    if best is None:
        raise RuntimeError("S1 all deterministic starts failed")
    if float(best.fun) + 1e-9 >= poiss_nll:
        return {
            "family": "nb2_iid",
            "mu": mu0,
            "kappa": 0.0,
            "poisson_boundary": True,
            "converged": True,
            "optimizer": "poisson boundary selected",
        }
    return {
        "family": "nb2_iid",
        "mu": math.exp(float(best.x[0])),
        "kappa": math.exp(float(best.x[1])),
        "poisson_boundary": math.exp(float(best.x[1])) < 1e-8,
        "converged": True,
        "optimizer": str(best.message),
    }


INGARCH_STARTS = ((0.10, 0.50), (0.30, 0.30), (0.01, 0.01))


def fit_ingarch(history: np.ndarray, negative_binomial: bool) -> dict:
    history = np.asarray(history, dtype=float)
    mean = max(float(history.mean()), 1e-8)
    starts = [ingarch_start(mean, a, b) for a, b in INGARCH_STARTS]
    if negative_binomial:
        s1 = fit_s1(history)
        kappa_starts = [max(float(s1.get("kappa", 0.0)), 1e-4), 0.1, 1.0]
    else:
        kappa_starts = [None]

    best = None
    for base in starts:
        for kap0 in kappa_starts:
            x0 = base if kap0 is None else np.r_[base, math.log(min(max(float(kap0), math.exp(-20)), math.exp(6)))]

            def obj(z):
                omega, a, b = ingarch_transform(z[:3])
                lam = ingarch_lambda(history, omega, a, b)
                if negative_binomial:
                    kap = math.exp(float(z[3]))
                    ll = nb2_logpmf(history, lam, kap)
                else:
                    ll = poisson_logpmf(history, lam)
                value = float(-ll.sum())
                return value if math.isfinite(value) else 1e100

            bounds = [(-20, 20), (-12, 12), (-12, 12)] + ([(-20, 6)] if negative_binomial else [])
            r = minimize(obj, x0, method="L-BFGS-B", bounds=bounds, options={"maxiter": 500, "ftol": 1e-10})
            if r.success and np.isfinite(r.fun) and (best is None or r.fun < best.fun):
                best = r
    if best is None:
        raise RuntimeError(("D1" if negative_binomial else "D0") + " all deterministic starts failed")
    omega, a, b = ingarch_transform(best.x[:3])
    kappa = math.exp(float(best.x[3])) if negative_binomial else 0.0
    return {
        "family": "nb2_ingarch11" if negative_binomial else "poisson_ingarch11",
        "omega": omega,
        "a": a,
        "b": b,
        "persistence": a + b,
        "kappa": kappa,
        "poisson_boundary": bool(negative_binomial and kappa < 1e-8),
        "optimizer": str(best.message),
        "converged": True,
    }


def forecast_model(history: np.ndarray, model: dict) -> tuple[float, float]:
    fam = model["family"]
    if fam in {"poisson_iid", "nb2_iid"}:
        return float(model["mu"]), float(model.get("kappa", 0.0))
    mu = ingarch_forecast(history, float(model["omega"]), float(model["a"]), float(model["b"]))
    return mu, float(model.get("kappa", 0.0))


def pmf_log_one(obs: float, mu: float, kappa: float) -> float:
    arr = np.array([obs], dtype=float)
    mm = np.array([mu], dtype=float)
    return float((nb2_logpmf(arr, mm, kappa) if kappa > 1e-10 else poisson_logpmf(arr, mm))[0])


def score_start_index(pairs) -> int:
    try:
        return pairs.index(SCORE_START)
    except ValueError as e:
        raise RuntimeError("FAIL_CLOSED SCORE_START absent") from e


def history_indices_for_origin(origin: int) -> np.ndarray:
    if origin <= 0:
        return np.array([], dtype=int)
    return np.arange(0, origin, dtype=int)


def moving_block_bootstrap(values: np.ndarray, block: int, reps: int, seed: int) -> dict:
    x = np.asarray(values, dtype=float)
    n = len(x)
    if n == 0 or block <= 0 or block > n:
        raise ValueError("invalid moving-block bootstrap inputs")
    starts = np.arange(0, n - block + 1, dtype=int)
    need = math.ceil(n / block)
    rng = np.random.default_rng(seed)
    means = np.empty(reps, dtype=float)
    for i in range(reps):
        chosen = rng.choice(starts, size=need, replace=True)
        sample = np.concatenate([x[s : s + block] for s in chosen])[:n]
        means[i] = float(sample.mean())
    return {
        "mean_improvement": float(x.mean()),
        "q025": float(np.quantile(means, 0.025)),
        "median": float(np.quantile(means, 0.5)),
        "q975": float(np.quantile(means, 0.975)),
        "block_length_months": int(block),
        "reps": int(reps),
        "seed": int(seed),
    }


def scipy_dist(mu: float, kappa: float):
    if kappa <= 1e-10:
        return poisson(mu)
    r = 1.0 / kappa
    p = r / (r + mu)
    return nbinom(r, p)


def diagnostics(obs: np.ndarray, mu: np.ndarray, kappa: np.ndarray, seed: int) -> dict:
    obs = np.asarray(obs, dtype=float)
    mu = np.asarray(mu, dtype=float)
    kappa = np.asarray(kappa, dtype=float)
    rng = np.random.default_rng(seed)
    u = rng.random(len(obs))
    pits = np.empty(len(obs), dtype=float)
    cover = {"50": 0, "90": 0, "95": 0}
    qs = {"50": (0.25, 0.75), "90": (0.05, 0.95), "95": (0.025, 0.975)}
    for i, (yy, mm, kk) in enumerate(zip(obs.astype(int), mu, kappa)):
        d = scipy_dist(float(mm), float(kk))
        prev = 0.0 if yy <= 0 else float(d.cdf(yy - 1))
        pits[i] = prev + u[i] * float(d.pmf(yy))
        for name, (loq, hiq) in qs.items():
            if int(d.ppf(loq)) <= yy <= int(d.ppf(hiq)):
                cover[name] += 1
    var = mu + kappa * mu * mu
    resid = (obs - mu) / np.sqrt(np.clip(var, 1e-12, None))

    def acf(lag: int):
        if len(resid) <= lag:
            return None
        x = resid[:-lag] - resid[:-lag].mean()
        z = resid[lag:] - resid[lag:].mean()
        den = math.sqrt(float(np.dot(x, x) * np.dot(z, z)))
        return None if den == 0 else float(np.dot(x, z) / den)

    rhos = [acf(i) for i in range(1, 13)]
    if all(v is not None and math.isfinite(float(v)) for v in rhos) and len(resid) > 12:
        q = len(resid) * (len(resid) + 2.0) * sum(float(rhos[k - 1]) ** 2 / (len(resid) - k) for k in range(1, 13))
        lb = {"Q12": float(q), "pvalue_chi2_df12": float(chi2.sf(q, 12))}
    else:
        lb = None
    ks = kstest(pits, "uniform")
    return {
        "randomized_pit": {"seed": seed, "mean": float(pits.mean()), "ks_statistic": float(ks.statistic), "ks_pvalue": float(ks.pvalue)},
        "central_interval_coverage": {k: cover[k] / len(obs) for k in cover},
        "pearson_residual": {
            "mean": float(resid.mean()),
            "variance_population": float(resid.var()),
            "acf_lag1": acf(1),
            "acf_lag12": acf(12),
            "ljung_box": lb,
        },
    }


def summarize_params(rows: list[dict], keys: tuple[str, ...]) -> dict:
    out = {}
    for key in keys:
        vals = np.array([float(r[key]) for r in rows if key in r and r[key] is not None and math.isfinite(float(r[key]))], dtype=float)
        if len(vals):
            out[key] = {
                "n": int(len(vals)),
                "median": float(np.median(vals)),
                "q025": float(np.quantile(vals, 0.025)),
                "q975": float(np.quantile(vals, 0.975)),
                "last": float(vals[-1]),
            }
    return out


def run_prequential(y: np.ndarray, pairs) -> dict:
    start = score_start_index(pairs)
    if start != 199:
        raise RuntimeError(f"FAIL_CLOSED warm-up months {start} != 199")
    origins = np.arange(start, len(y), dtype=int)
    records = {m: {"nll": [], "mu": [], "kappa": [], "params": []} for m in MODEL_NAMES}
    observations = []
    scored_pairs = []

    for origin in origins:
        hist_idx = history_indices_for_origin(int(origin))
        if len(hist_idx) != origin or (len(hist_idx) and hist_idx[-1] != origin - 1):
            raise RuntimeError("FAIL_CLOSED prequential history firewall")
        history = y[hist_idx]
        obs = float(y[origin])
        fitted = {
            "S0": fit_s0(history),
            "S1": fit_s1(history),
            "D0": fit_ingarch(history, negative_binomial=False),
            "D1": fit_ingarch(history, negative_binomial=True),
        }
        observations.append(obs)
        scored_pairs.append(pairs[origin])
        for name, model in fitted.items():
            mu, kap = forecast_model(history, model)
            ll = pmf_log_one(obs, mu, kap)
            if not math.isfinite(ll):
                raise RuntimeError(f"{name} non-finite one-step log score at {pairs[origin]}")
            records[name]["nll"].append(-ll)
            records[name]["mu"].append(mu)
            records[name]["kappa"].append(kap)
            records[name]["params"].append(model)

    obs = np.asarray(observations, dtype=float)
    result_models = {}
    for name in MODEL_NAMES:
        nll = np.asarray(records[name]["nll"], dtype=float)
        mu = np.asarray(records[name]["mu"], dtype=float)
        kap = np.asarray(records[name]["kappa"], dtype=float)
        if name in {"D0", "D1"}:
            ps = summarize_params(records[name]["params"], ("omega", "a", "b", "persistence", "kappa"))
        else:
            ps = summarize_params(records[name]["params"], ("mu", "kappa"))
        result_models[name] = {
            "mean_nll": float(nll.mean()),
            "diagnostics": diagnostics(obs, mu, kap, SEED),
            "parameter_trajectory_summary": ps,
            "fit_failures": 0,
        }

    contrasts = {}
    for label, comp, chall in [
        ("S0-D0", "S0", "D0"),
        ("S1-D1", "S1", "D1"),
        ("D0-D1", "D0", "D1"),
        ("S0-S1", "S0", "S1"),
    ]:
        diff = np.asarray(records[comp]["nll"]) - np.asarray(records[chall]["nll"])
        contrasts[label] = {
            "primary_12m": moving_block_bootstrap(diff, PRIMARY_BLOCK, BOOT_REPS, SEED),
            "secondary_6m": moving_block_bootstrap(diff, SECONDARY_BLOCK, BOOT_REPS, SEED),
        }

    def stratum(lo_pair, hi_pair):
        mask = np.array([lo_pair <= p <= hi_pair for p in scored_pairs], dtype=bool)
        sub = {}
        for name in MODEL_NAMES:
            nll = np.asarray(records[name]["nll"])[mask]
            mu = np.asarray(records[name]["mu"])[mask]
            kap = np.asarray(records[name]["kappa"])[mask]
            sub[name] = {
                "months": int(mask.sum()),
                "mean_nll": float(nll.mean()),
                "diagnostics": diagnostics(obs[mask], mu, kap, SEED),
            }
        return sub

    return {
        "scored_months": int(len(origins)),
        "scored_start": f"{scored_pairs[0][0]:04d}-{scored_pairs[0][1]:02d}",
        "scored_end": f"{scored_pairs[-1][0]:04d}-{scored_pairs[-1][1]:02d}",
        "models": result_models,
        "contrasts": contrasts,
        "strata": {
            "early_2005_2014": stratum((2005, 1), (2014, 12)),
            "previously_opened_2015_2022_11": stratum((2015, 1), (2022, 11)),
        },
        "raw_scored_count_mean": float(obs.mean()),
        "raw_scored_count_variance_population": float(obs.var()),
    }


def main() -> int:
    data = fetch_workbook()
    valid, excluded = parse_exact_dates(data)
    keys, pairs, y = monthly_series(valid)
    result = run_prequential(y, pairs)
    receipt = {
        "schema": "pcold-ingarch-prequential-v0",
        "status": "RETROSPECTIVE_PREQUENTIAL_EXECUTED_NO_UNTOUCHED_HOLDOUT_CLAIM",
        "ffbk_preregistration_pr": 1501,
        "source": {
            "doi": "10.6084/m9.figshare.23940069.v3",
            "article_id": ARTICLE_ID,
            "version": VERSION,
            "file_id": FILE_ID,
            "filename": WORKBOOK_NAME,
            "sha256": WORKBOOK_SHA256,
            "valid_exact_occurrence_month_rows": len(valid),
            "excluded_occurrence_date_rows": len(excluded),
            "monthly_series_start": f"{EXPECTED_MIN_MONTH[0]:04d}-{EXPECTED_MIN_MONTH[1]:02d}",
            "monthly_series_end": f"{EXPECTED_MAX_MONTH[0]:04d}-{EXPECTED_MAX_MONTH[1]:02d}",
        },
        "execution": {
            "python": platform.python_version(),
            "numpy": np.__version__,
            "scipy": scipy.__version__,
            "openpyxl": openpyxl.__version__,
            "bootstrap_seed": SEED,
            "bootstrap_reps": BOOT_REPS,
            "primary_block_length_months": PRIMARY_BLOCK,
            "secondary_block_length_months": SECONDARY_BLOCK,
            "persistence_cap": PERSISTENCE_CAP,
            "ingarch_start_grid": [list(x) for x in INGARCH_STARTS],
        },
        "prequential": result,
        "claim_boundary": [
            "retrospective/prequential benchmark after prior P-COLD target access",
            "no untouched holdout claim",
            "no insurer/P&C transfer",
            "no production/default family promotion",
            "no NextGen API promotion",
            "no tail or capital quantity produced in this run",
        ],
    }
    out = Path("research-results/pcold-ingarch-prequential-v0.json")
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(receipt, indent=2, sort_keys=True) + "\n")
    print(json.dumps(receipt, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
